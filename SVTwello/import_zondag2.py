from pathlib import Path
import json
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent
EXCEL_PATH = BASE / 'zondag2.xlsx'


def parse_bool(value):
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    return str(value).strip().lower() in {'ja', 'yes', 'true', '1', 'x'}


def parse_truthy(value):
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    return str(value).strip().lower() in {'ja', 'yes', 'true', '1', 'x', 'y'}


def normalize_foot(value):
    if value is None:
        return 'rechts'
    text = str(value).strip().lower()
    if text in {'links', 'linker', 'left'}:
        return 'links'
    return 'rechts'


def load_sheet_rows(sheet_name):
    wb = load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else '' for h in rows[0]]
    data = []
    for row in rows[1:]:
        if not any(v is not None and str(v).strip() != '' for v in row):
            continue
        item = {}
        for idx, header in enumerate(headers):
            value = row[idx] if idx < len(row) else None
            item[header] = value
        data.append(item)
    return data


def build_matches_payload():
    wedstrijden_rows = load_sheet_rows('wedstrijden')
    invoer_rows = load_sheet_rows('wedstrijdinvoer')

    by_match = {}
    for row in invoer_rows:
        match_id = row.get('wedstrijd_id')
        if not match_id:
            continue
        entry = by_match.setdefault(match_id, {
            'doelpunten': [],
            'assists': [],
            'penalties': [],
            'kaarten': [],
            'statussen': [],
            'te_laat': []
        })

        speler_naam = row.get('speler_naam') or ''
        doelpunten = row.get('doelpunten')
        assists = row.get('assists')
        penalty = row.get('penalty')
        geel = row.get('geel')
        rood = row.get('rood')
        te_laat = row.get('te laat')
        status = row.get('status')

        if doelpunten not in (None, '', 0):
            count = int(doelpunten)
            for _ in range(count):
                entry['doelpunten'].append({'speler': speler_naam})

        if assists not in (None, '', 0):
            count = int(assists)
            for _ in range(count):
                entry['assists'].append({'speler': speler_naam})

        if penalty not in (None, '', 0):
            count = int(penalty)
            for _ in range(count):
                entry['penalties'].append({'speler': speler_naam})

        if geel not in (None, '', 0):
            entry['kaarten'].append({'speler': speler_naam, 'type': 'geel'})
        if rood not in (None, '', 0):
            entry['kaarten'].append({'speler': speler_naam, 'type': 'rood'})
        if te_laat not in (None, '', 0):
            entry['te_laat'].append({'speler': speler_naam, 'waarde': te_laat})
        if status not in (None, '', 0):
            entry['statussen'].append({'speler': speler_naam, 'status': status})

    wedstrijden = []
    for row in wedstrijden_rows:
        match_id = row.get('wedstrijd_id')
        if not match_id:
            continue
        data = by_match.get(match_id, {})
        wedstrijden.append({
            'id': match_id,
            'datum': row.get('datum'),
            'tijd': row.get('tijd'),
            'thuis': row.get('thuis'),
            'uit': row.get('uit'),
            'uitslag': row.get('uitslag') or '',
            'competitie': row.get('competitie') or 'Competitie',
            'doelpunten': data.get('doelpunten', []),
            'assists': data.get('assists', []),
            'penalties': data.get('penalties', []),
            'kaarten': data.get('kaarten', []),
            'te_laat': data.get('te_laat', []),
            'statussen': data.get('statussen', []),
        })

    return {
        'team': 'SV Twello 2',
        'seizoen': '2026/2027',
        'wedstrijden': wedstrijden,
    }


def build_players_json():
    spelers_rows = load_sheet_rows('spelers')
    staf_rows = load_sheet_rows('staf')
    matches_payload = build_matches_payload()
    trainings_rows = load_sheet_rows('trainingsinvoer')
    wedstrijd_rows = load_sheet_rows('wedstrijdinvoer')

    spelers = []
    for row in spelers_rows:
        naam = row.get('naam')
        if not naam:
            continue
        spelers.append({
            'rugnummer': row.get('rugnummer'),
            'naam': naam,
            'positie': row.get('positie') or '',
            'status': row.get('Status') or '',
            'voet': normalize_foot(row.get('voet')),
            'gastspeler': parse_bool(row.get('gastspeler')),
            'aanvoerder': parse_bool(row.get('aanvoerder')),
        })

    player_lookup = {str(player['naam']).strip().lower(): player for player in spelers if player.get('naam')}

    if trainings_rows:
        training_headers = [h for h in trainings_rows[0].keys() if h not in {'speler_id', 'speler naam'}]
        training_total = len(training_headers)
        for row in trainings_rows[1:]:
            player_name = str(row.get('speler naam') or '').strip()
            key = player_name.lower()
            player = player_lookup.get(key)
            if not player:
                continue
            attended = 0
            for header in training_headers:
                value = row.get(header)
                if parse_truthy(value):
                    attended += 1
            player['training'] = attended
            player['trainingTotal'] = training_total

    if wedstrijd_rows:
        attendance_totals = {}
        attendance_counts = {}
        for row in wedstrijd_rows:
            player_name = str(row.get('speler_naam') or '').strip()
            if not player_name:
                continue
            key = player_name.lower()
            attendance_totals[key] = attendance_totals.get(key, 0) + 1
            if parse_truthy(row.get('aanwezig')):
                attendance_counts[key] = attendance_counts.get(key, 0) + 1

        for key, player in player_lookup.items():
            player['attendance'] = attendance_counts.get(key, 0)
            player['attendanceTotal'] = attendance_totals.get(key, 0)

    staf = []
    for row in staf_rows:
        naam = row.get('naam')
        if not naam:
            continue
        staf.append({
            'naam': naam,
            'rol': row.get('rol') or '',
        })

    payload = {
        'spelers': spelers,
        'staf': staf,
        'team': matches_payload['team'],
        'seizoen': matches_payload['seizoen'],
        'wedstrijden': matches_payload['wedstrijden'],
    }
    (BASE / 'spelers.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'Wrote {len(spelers)} players, {len(staf)} staff members and {len(matches_payload["wedstrijden"])} matches to spelers.json')


if __name__ == '__main__':
    build_players_json()
